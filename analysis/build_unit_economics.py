"""
build_unit_economics.py
-----------------------
Builds model/unit_economics.xlsx: customer lifetime value (LTV) vs acquisition
cost (CAC), and the break-even test for the proposed loyalty tier.

Every calculation in the workbook is a live Excel formula, so a reader can
change any input and the model recalculates. Inputs are colour-coded:
  blue text            = sourced input (Swiggy Q1 FY2027 shareholder letter)
  blue on yellow fill  = assumption the reader should challenge

Sourced inputs (Swiggy Q1 FY2027 shareholder letter, quarter ended 30 June 2026):
  Food delivery GOV INR 9,490 Cr (page 10); food delivery average monthly
  transacting users 19.2 Mn and Adjusted EBITDA INR 292 Cr (page 6);
  medium-term Adjusted EBITDA guidance of 5% of GOV (page 10).
  https://www.swiggy.com/corporate/wp-content/uploads/2026/07/Q1-FY2027-Shareholder-letter.pdf

Run:
  python analysis/build_unit_economics.py
then recalculate (Excel does this on open; LibreOffice: see README).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE = Path(__file__).resolve().parent.parent
OUT = BASE / "model" / "unit_economics.xlsx"
OUT.parent.mkdir(parents=True, exist_ok=True)

SOURCE = "Swiggy Q1 FY2027 shareholder letter (quarter ended 30 Jun 2026)"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
YELLOW = PatternFill("solid", fgColor="FFFF00")
SECTION = PatternFill("solid", fgColor="E8EEF5")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
INR = '"₹"#,##0;("₹"#,##0);-'
INR1 = '"₹"#,##0.0;("₹"#,##0.0);-'
PCT1 = "0.0%;(0.0%);-"
PCT2 = "0.00%;(0.00%);-"
MULT = '0.0"x";(0.0"x");-'
MONTHS = '0.0" mo";(0.0" mo");-'


def label(ws, cell: str, text: str, font: Font = BLACK) -> None:
    ws[cell] = text
    ws[cell].font = font


def value(ws, cell: str, v, fmt: str, font: Font, note: str | None = None, assumption: bool = False) -> None:
    ws[cell] = v
    ws[cell].font = font
    ws[cell].number_format = fmt
    if assumption:
        ws[cell].fill = YELLOW
    if note:
        ws[cell].comment = Comment(note, "Model")


def section(ws, row: int, text: str, width: int = 7) -> None:
    for col in range(1, width + 1):
        ws.cell(row=row, column=col).fill = SECTION
    ws.cell(row=row, column=1, value=text).font = BOLD


def build() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    ws.column_dimensions["A"].width = 52
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 15

    label(ws, "A1", "Swiggy food delivery: customer unit economics and loyalty-tier break-even", TITLE)
    label(ws, "A2", "Blue = sourced input · Blue on yellow = assumption to challenge · Black = formula. "
                    "Change any blue cell; everything else recalculates.")
    ws["A2"].font = Font(name=FONT, italic=True, color="595959")

    # ---------------------------------------------------------------- inputs
    section(ws, 4, "1. Sourced inputs (" + SOURCE + ")")
    label(ws, "A5", "Food delivery GOV in the quarter (₹ crore)")
    value(ws, "B5", 9490, '#,##0', BLUE, f"{SOURCE}, page 10: food delivery GOV grew 17.4% YoY to INR 9,490 Cr.")
    label(ws, "A6", "Food delivery average monthly transacting users (million)")
    value(ws, "B6", 19.2, '0.0', BLUE, f"{SOURCE}, page 6, Food delivery table: Average MTU 19.2 million in Q1FY27.")
    label(ws, "A7", "Food delivery Adjusted EBITDA in the quarter (₹ crore)")
    value(ws, "B7", 292, '#,##0', BLUE, f"{SOURCE}, page 6, Food delivery table: Adjusted EBITDA INR 292 Cr in Q1FY27.")
    label(ws, "A8", "Medium-term Adjusted EBITDA guidance (% of GOV)")
    value(ws, "B8", 0.05, PCT1, BLUE, f"{SOURCE}, page 10: steady-state guidance of 5% Adjusted EBITDA on GOV.")
    label(ws, "A9", "Months in the quarter")
    value(ws, "B9", 3, '0', BLUE)

    section(ws, 11, "2. Derived per-user economics")
    label(ws, "A12", "GOV per active user per month (₹)")
    ws["B12"] = "=B5*10^7/(B9*B6*10^6)"
    ws["B12"].number_format = INR1
    label(ws, "A13", "Reported Adjusted EBITDA margin (% of GOV)")
    ws["B13"] = "=B7/B5"
    ws["B13"].number_format = PCT2
    label(ws, "A14", "Monthly profit per active user at reported margin (₹)")
    ws["B14"] = "=B12*B13"
    ws["B14"].number_format = INR1
    label(ws, "A15", "Monthly profit per active user at guided margin (₹)")
    ws["B15"] = "=B12*B8"
    ws["B15"].number_format = INR1
    label(ws, "A16", "Note: Adjusted EBITDA is after fixed costs, so it understates per-order contribution. "
                     "Contribution margin is not stated in the letter's text; this model is deliberately conservative.")
    ws["A16"].font = Font(name=FONT, italic=True, size=9, color="595959")

    # ----------------------------------------------------------- assumptions
    section(ws, 18, "3. Assumptions (not published; challenge these)")
    label(ws, "A19", "Customer acquisition cost, CAC (₹ per new user)")
    value(ws, "B19", 400, INR, BLUE, "Assumption. Mid-point of the ₹300-500 range used in the case study; "
                                      "Swiggy does not publish CAC.", assumption=True)
    label(ws, "A20", "Monthly churn of an active user (base case)")
    value(ws, "B20", 0.10, PCT1, BLUE, "Assumption. Probability an active user stops ordering in a given month. "
                                        "Swiggy does not publish churn.", assumption=True)
    label(ws, "A21", "Annual discount rate")
    value(ws, "B21", 0.12, PCT1, BLUE, "Assumption: cost of capital used to discount future profit.", assumption=True)
    label(ws, "A22", "Loyalty-tier benefit cost per member per month (₹)")
    value(ws, "B22", 25, INR, BLUE, "Assumption: e.g. delivery-fee waivers and coins redeemed per member per month.",
          assumption=True)

    # --------------------------------------------------------------- LTV/CAC
    section(ws, 24, "4. Base case: lifetime value vs acquisition cost")
    ws["B25"], ws["C25"] = "Reported margin", "Guided margin"
    for c in ("B25", "C25"):
        ws[c].font = HEAD
        ws[c].fill = HEAD_FILL
        ws[c].alignment = Alignment(horizontal="center")
    label(ws, "A26", "Monthly discount rate")
    ws["B26"] = "=(1+$B$21)^(1/12)-1"
    ws["C26"] = "=(1+$B$21)^(1/12)-1"
    label(ws, "A27", "Expected discounted lifetime (months)")
    ws["B27"] = "=(1+B26)/(B26+$B$20)"
    ws["C27"] = "=(1+C26)/(C26+$B$20)"
    label(ws, "A28", "Lifetime value, LTV (₹)")
    ws["B28"] = "=$B$14*B27"
    ws["C28"] = "=$B$15*C27"
    label(ws, "A29", "LTV / CAC")
    ws["B29"] = "=IFERROR(B28/$B$19,0)"
    ws["C29"] = "=IFERROR(C28/$B$19,0)"
    label(ws, "A30", "Simple CAC payback (months of profit)")
    ws["B30"] = "=IFERROR($B$19/$B$14,0)"
    ws["C30"] = "=IFERROR($B$19/$B$15,0)"
    label(ws, "A31", "Maximum affordable CAC at LTV/CAC = 3x (₹)")
    ws["B31"] = "=B28/3"
    ws["C31"] = "=C28/3"
    for r, fmt in ((26, PCT2), (27, MONTHS), (28, INR), (29, MULT), (30, MONTHS), (31, INR)):
        ws[f"B{r}"].number_format = fmt
        ws[f"C{r}"].number_format = fmt
    label(ws, "A32", "LTV formula: monthly profit × Σ (retention / (1+discount))^t = monthly profit × (1+d)/(d+churn).")
    ws["A32"].font = Font(name=FONT, italic=True, size=9, color="595959")

    # ------------------------------------------------------- sensitivity grid
    section(ws, 34, "5. Sensitivity: LTV / CAC at the reported margin (rows: monthly churn · columns: CAC ₹)")
    churns = [0.06, 0.08, 0.10, 0.12, 0.15, 0.20]
    cacs = [200, 300, 400, 500, 700]
    ws["A35"] = "Monthly churn ↓   CAC →"
    ws["A35"].font = HEAD
    ws["A35"].fill = HEAD_FILL
    for j, cac in enumerate(cacs):
        cell = ws.cell(row=35, column=2 + j, value=cac)
        cell.font, cell.fill, cell.number_format = HEAD, HEAD_FILL, INR
    for i, churn in enumerate(churns):
        r = 36 + i
        c = ws.cell(row=r, column=1, value=churn)
        c.font, c.number_format, c.alignment = BLUE, PCT1, Alignment(horizontal="right")
        for j in range(len(cacs)):
            col = chr(ord("B") + j)
            ws[f"{col}{r}"] = f"=$B$14*(1+$B$26)/($B$26+$A{r})/{col}$35"
            ws[f"{col}{r}"].number_format = MULT

    section(ws, 43, "6. Same grid at the guided 5% margin")
    ws["A44"] = "Monthly churn ↓   CAC →"
    ws["A44"].font = HEAD
    ws["A44"].fill = HEAD_FILL
    for j, cac in enumerate(cacs):
        cell = ws.cell(row=44, column=2 + j, value=cac)
        cell.font, cell.fill, cell.number_format = HEAD, HEAD_FILL, INR
    for i, churn in enumerate(churns):
        r = 45 + i
        c = ws.cell(row=r, column=1, value=churn)
        c.font, c.number_format, c.alignment = BLUE, PCT1, Alignment(horizontal="right")
        for j in range(len(cacs)):
            col = chr(ord("B") + j)
            ws[f"{col}{r}"] = f"=$B$15*(1+$C$26)/($C$26+$A{r})/{col}$44"
            ws[f"{col}{r}"].number_format = MULT

    # -------------------------------------------------- loyalty break-even
    section(ws, 52, "7. Loyalty tier: what must it achieve to pay for itself?")
    ws["B53"], ws["C53"] = "Reported margin", "Guided margin"
    for c in ("B53", "C53"):
        ws[c].font = HEAD
        ws[c].fill = HEAD_FILL
        ws[c].alignment = Alignment(horizontal="center")
    label(ws, "A54", "Monthly profit per active user (₹)")
    ws["B54"] = "=B14"
    ws["C54"] = "=B15"
    label(ws, "A55", "Benefit cost as % of monthly profit per user")
    ws["B55"] = "=IFERROR($B$22/B54,0)"
    ws["C55"] = "=IFERROR($B$22/C54,0)"
    label(ws, "A56", "Route A: extra GOV per member per month needed (₹)")
    ws["B56"] = "=$B$22/$B$13"
    ws["C56"] = "=$B$22/$B$8"
    label(ws, "A57", "Route A as % increase in spend per member")
    ws["B57"] = "=B56/$B$12"
    ws["C57"] = "=C56/$B$12"
    label(ws, "A58", "Route B: maximum monthly churn with the tier (break-even)")
    ws["B58"] = '=IF(B54<=$B$22,"Not possible",(B54-$B$22)*($B$26+$B$20)/B54-$B$26)'
    ws["C58"] = '=IF(C54<=$B$22,"Not possible",(C54-$B$22)*($C$26+$B$20)/C54-$C$26)'
    label(ws, "A59", "Route B: churn reduction needed (percentage points)")
    ws["B59"] = '=IF(ISNUMBER(B58),$B$20-B58,"Not possible")'
    ws["C59"] = '=IF(ISNUMBER(C58),$B$20-C58,"Not possible")'
    for r, fmt in ((54, INR1), (55, PCT1), (56, INR), (57, PCT1), (58, PCT2), (59, PCT2)):
        ws[f"B{r}"].number_format = fmt
        ws[f"C{r}"].number_format = fmt
    label(ws, "A60", "Route A pays through extra spend alone; Route B pays through members staying longer. "
                     "Route B assumes members keep the same spend and that the benefit cost applies to every member.")
    ws["A60"].font = Font(name=FONT, italic=True, size=9, color="595959")

    section(ws, 62, "8. Route B by benefit cost (reported margin): churn reduction needed, percentage points")
    ws["A63"] = "Benefit cost per member per month (₹) →"
    ws["A63"].font = HEAD
    ws["A63"].fill = HEAD_FILL
    for j, cost in enumerate([10, 15, 25, 35, 45]):
        col = chr(ord("B") + j)
        cell = ws[f"{col}63"]
        cell.value, cell.font, cell.fill, cell.number_format = cost, HEAD, HEAD_FILL, INR
        ws[f"{col}64"] = (f'=IF($B$14<={col}63,"Not possible",'
                          f'$B$20-(($B$14-{col}63)*($B$26+$B$20)/$B$14-$B$26))')
        ws[f"{col}64"].number_format = PCT2
    label(ws, "A64", "Churn reduction needed (base churn in B20)")

    for row in ws.iter_rows(min_row=5, max_row=64):
        for cell in row:
            if cell.font is None or cell.font.name != FONT:
                cell.font = Font(name=FONT, bold=cell.font.bold if cell.font else False,
                                 color=cell.font.color if cell.font else None)
    ws.freeze_panes = "B4"

    # ------------------------------------------------------------- sources
    src = wb.create_sheet("Sources")
    src.column_dimensions["A"].width = 30
    src.column_dimensions["B"].width = 110
    rows = [
        ("Item", "Source / basis"),
        ("Food delivery GOV ₹9,490 Cr", f"{SOURCE}, page 10 (Management Perspectives, Q1)."),
        ("Food delivery MTU 19.2 Mn", f"{SOURCE}, page 6, Food delivery operating table."),
        ("Food delivery Adj. EBITDA ₹292 Cr", f"{SOURCE}, page 6, Food delivery table; 3.1% of GOV stated on page 2."),
        ("5% Adj. EBITDA guidance", f"{SOURCE}, page 10: steady-state guidance over the medium term."),
        ("Letter URL", "https://www.swiggy.com/corporate/wp-content/uploads/2026/07/Q1-FY2027-Shareholder-letter.pdf"),
        ("CAC, churn, discount rate, benefit cost", "Assumptions. Not published by Swiggy. Shown in yellow; test them with the grids."),
        ("Why EBITDA, not contribution", "Contribution margin is not stated in the letter's text. Using Adjusted EBITDA "
                                         "understates profit per user, so the model errs on the cautious side."),
    ]
    for i, (a, b) in enumerate(rows, start=1):
        src.cell(row=i, column=1, value=a).font = BOLD if i == 1 else BLACK
        src.cell(row=i, column=2, value=b).font = BOLD if i == 1 else BLACK
        src.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

    wb.save(OUT)
    print(f"Wrote {OUT.relative_to(BASE)}")


if __name__ == "__main__":
    build()
